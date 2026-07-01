"""
Sistema de Cronogramas de Actividades - INEB Domingo Juarros
Stack: Python 3 + Flask + SQLite + Jinja2
"""

from flask import Flask, render_template, redirect, url_for, request, flash, jsonify, abort
from flask_sqlalchemy import SQLAlchemy
from flask_login import (LoginManager, UserMixin, login_user,
                          logout_user, login_required, current_user)
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime
import json, os

# ─────────────────────────────────────────────
# APP CONFIG
# ─────────────────────────────────────────────
app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'cambiar-en-produccion-2026')

# Ruta de BD compatible con Render y local
BASE_DIR = os.path.abspath(os.path.dirname(__file__))
DB_DIR   = '/tmp' if os.environ.get('RENDER') else BASE_DIR
DB_PATH  = os.path.join(DB_DIR, 'cronogramas.db')
app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{DB_PATH}'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Debes iniciar sesión para acceder.'
login_manager.login_message_category = 'warning'

# ─────────────────────────────────────────────
# MODELOS (Base de Datos)
# ─────────────────────────────────────────────

class Maestro(UserMixin, db.Model):
    """Tabla de maestros con credenciales"""
    __tablename__ = 'maestros'
    id            = db.Column(db.Integer, primary_key=True)
    nombre        = db.Column(db.String(120), nullable=False)
    email         = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    creado_en     = db.Column(db.DateTime, default=datetime.utcnow)
    cronogramas   = db.relationship('Cronograma', backref='maestro', lazy=True,
                                     cascade='all, delete-orphan')

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

    def __repr__(self):
        return f'<Maestro {self.nombre}>'


class Cronograma(db.Model):
    """Cronograma bimestral de actividades"""
    __tablename__ = 'cronogramas'
    id              = db.Column(db.Integer, primary_key=True)
    maestro_id      = db.Column(db.Integer, db.ForeignKey('maestros.id'), nullable=False)
    catedra         = db.Column(db.String(120), nullable=False)
    grado           = db.Column(db.String(80),  nullable=False)
    seccion         = db.Column(db.String(10),  nullable=False)
    unidad          = db.Column(db.String(80),  nullable=False)
    ciclo           = db.Column(db.String(10),  nullable=False)
    fecha_inicio    = db.Column(db.String(20),  nullable=False)
    fecha_fin       = db.Column(db.String(20),  nullable=False)
    punteo_meta     = db.Column(db.Integer, default=100)
    creado_en       = db.Column(db.DateTime, default=datetime.utcnow)
    actualizado_en  = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    actividades     = db.relationship('Actividad', backref='cronograma', lazy=True,
                                       cascade='all, delete-orphan',
                                       order_by='Actividad.numero')

    @property
    def total_unidad(self):
        return sum(a.total_actividad for a in self.actividades)

    def __repr__(self):
        return f'<Cronograma {self.catedra} - {self.grado} {self.seccion}>'


class Actividad(db.Model):
    """Actividad dentro de un cronograma (máx 6 por cronograma)"""
    __tablename__ = 'actividades'
    id              = db.Column(db.Integer, primary_key=True)
    cronograma_id   = db.Column(db.Integer, db.ForeignKey('cronogramas.id'), nullable=False)
    numero          = db.Column(db.Integer, nullable=False)   # 1-6
    descripcion     = db.Column(db.Text,    nullable=False)
    fecha_actividad = db.Column(db.String(20), nullable=True)
    aspectos_json   = db.Column(db.Text, default='[]')

    @property
    def aspectos(self):
        return json.loads(self.aspectos_json or '[]')

    @aspectos.setter
    def aspectos(self, value):
        self.aspectos_json = json.dumps(value, ensure_ascii=False)

    @property
    def total_actividad(self):
        return sum(int(a.get('puntaje', 0)) for a in self.aspectos)

    def __repr__(self):
        return f'<Actividad #{self.numero}>'


# ─────────────────────────────────────────────
# LOGIN MANAGER
# ─────────────────────────────────────────────

@login_manager.user_loader
def load_user(user_id):
    return db.session.get(Maestro, int(user_id))


# ─────────────────────────────────────────────
# RUTAS DE AUTENTICACIÓN
# ─────────────────────────────────────────────

@app.route('/')
def index():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        maestro  = Maestro.query.filter_by(email=email).first()
        if maestro and maestro.check_password(password):
            login_user(maestro, remember=True)
            next_page = request.args.get('next')
            return redirect(next_page or url_for('dashboard'))
        flash('Correo o contraseña incorrectos.', 'danger')
    return render_template('login.html')



@app.route('/recuperar', methods=['GET', 'POST'])
def recuperar():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        confirm  = request.form.get('confirm', '')
        maestro  = Maestro.query.filter_by(email=email).first()
        if not maestro:
            flash('No existe una cuenta con ese correo.', 'danger')
        elif password != confirm:
            flash('Las contrasenas no coinciden.', 'danger')
        elif len(password) < 6:
            flash('La contrasena debe tener al menos 6 caracteres.', 'danger')
        else:
            maestro.set_password(password)
            db.session.commit()
            flash('Contrasena actualizada correctamente. Ya puedes iniciar sesion.', 'success')
            return redirect(url_for('login'))
    return render_template('recuperar.html')

@app.route('/registro', methods=['GET', 'POST'])
def registro():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        nombre   = request.form.get('nombre', '').strip()
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        confirm  = request.form.get('confirm', '')
        if not nombre or not email or not password:
            flash('Todos los campos son requeridos.', 'danger')
        elif password != confirm:
            flash('Las contraseñas no coinciden.', 'danger')
        elif len(password) < 6:
            flash('La contraseña debe tener al menos 6 caracteres.', 'danger')
        elif Maestro.query.filter_by(email=email).first():
            flash('Ese correo ya está registrado.', 'danger')
        else:
            m = Maestro(nombre=nombre, email=email)
            m.set_password(password)
            db.session.add(m)
            db.session.commit()
            login_user(m)
            flash(f'Bienvenido/a, {nombre}!', 'success')
            return redirect(url_for('dashboard'))
    return render_template('registro.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))


# ─────────────────────────────────────────────
# DASHBOARD
# ─────────────────────────────────────────────

@app.route('/dashboard')
@login_required
def dashboard():
    cronogramas = (Cronograma.query
                   .filter_by(maestro_id=current_user.id)
                   .order_by(Cronograma.actualizado_en.desc())
                   .all())
    return render_template('dashboard.html', cronogramas=cronogramas)


# ─────────────────────────────────────────────
# CRUD CRONOGRAMAS
# ─────────────────────────────────────────────

GRADOS = ['Primero Básico', 'Segundo Básico', 'Tercero Básico']

@app.route('/cronograma/nuevo', methods=['GET', 'POST'])
@login_required
def nuevo_cronograma():
    if request.method == 'POST':
        cron = Cronograma(
            maestro_id   = current_user.id,
            catedra      = request.form.get('catedra', '').strip(),
            grado        = request.form.get('grado', ''),
            seccion      = request.form.get('seccion', '').strip().upper(),
            unidad       = request.form.get('unidad', '').strip(),
            ciclo        = request.form.get('ciclo', '').strip(),
            fecha_inicio = request.form.get('fecha_inicio', ''),
            fecha_fin    = request.form.get('fecha_fin', ''),
            punteo_meta  = int(request.form.get('punteo_meta', 100) or 100),
        )
        db.session.add(cron)
        db.session.flush()

        for i in range(1, 7):
            desc = request.form.get(f'act_{i}_descripcion', '').strip()
            if not desc:
                continue
            aspectos = []
            j = 1
            while True:
                nombre  = request.form.get(f'act_{i}_asp_{j}_nombre', '').strip()
                puntaje = request.form.get(f'act_{i}_asp_{j}_puntaje', '').strip()
                if not nombre:
                    break
                aspectos.append({'nombre': nombre, 'puntaje': int(puntaje or 0)})
                j += 1
            act = Actividad(
                cronograma_id   = cron.id,
                numero          = i,
                descripcion     = desc,
                fecha_actividad = request.form.get(f'act_{i}_fecha', ''),
            )
            act.aspectos = aspectos
            db.session.add(act)

        db.session.commit()
        flash('Cronograma guardado correctamente.', 'success')
        return redirect(url_for('dashboard'))

    return render_template('form_cronograma.html', grados=GRADOS,
                            cronograma=None, action='nuevo')


@app.route('/cronograma/<int:cid>/editar', methods=['GET', 'POST'])
@login_required
def editar_cronograma(cid):
    cron = Cronograma.query.get_or_404(cid)
    if cron.maestro_id != current_user.id:
        abort(403)

    if request.method == 'POST':
        cron.catedra      = request.form.get('catedra', '').strip()
        cron.grado        = request.form.get('grado', '')
        cron.seccion      = request.form.get('seccion', '').strip().upper()
        cron.unidad       = request.form.get('unidad', '').strip()
        cron.ciclo        = request.form.get('ciclo', '').strip()
        cron.fecha_inicio = request.form.get('fecha_inicio', '')
        cron.fecha_fin    = request.form.get('fecha_fin', '')
        cron.punteo_meta  = int(request.form.get('punteo_meta', 100) or 100)
        cron.actualizado_en = datetime.utcnow()

        for a in cron.actividades:
            db.session.delete(a)
        db.session.flush()

        for i in range(1, 7):
            desc = request.form.get(f'act_{i}_descripcion', '').strip()
            if not desc:
                continue
            aspectos = []
            j = 1
            while True:
                nombre  = request.form.get(f'act_{i}_asp_{j}_nombre', '').strip()
                puntaje = request.form.get(f'act_{i}_asp_{j}_puntaje', '').strip()
                if not nombre:
                    break
                aspectos.append({'nombre': nombre, 'puntaje': int(puntaje or 0)})
                j += 1
            act = Actividad(
                cronograma_id   = cron.id,
                numero          = i,
                descripcion     = desc,
                fecha_actividad = request.form.get(f'act_{i}_fecha', ''),
            )
            act.aspectos = aspectos
            db.session.add(act)

        db.session.commit()
        flash('Cronograma actualizado.', 'success')
        return redirect(url_for('dashboard'))

    return render_template('form_cronograma.html', grados=GRADOS,
                            cronograma=cron, action='editar')


@app.route('/cronograma/<int:cid>/eliminar', methods=['POST'])
@login_required
def eliminar_cronograma(cid):
    cron = Cronograma.query.get_or_404(cid)
    if cron.maestro_id != current_user.id:
        abort(403)
    db.session.delete(cron)
    db.session.commit()
    flash('Cronograma eliminado.', 'info')
    return redirect(url_for('dashboard'))


# ─────────────────────────────────────────────
# VISTA DE IMPRESIÓN
# ─────────────────────────────────────────────

@app.route('/imprimir')
@login_required
def imprimir():
    cronogramas = (Cronograma.query
                   .filter_by(maestro_id=current_user.id)
                   .order_by(Cronograma.grado, Cronograma.seccion)
                   .all())
    return render_template('imprimir.html', cronogramas=cronogramas,
                            maestro=current_user)


@app.route('/imprimir/<int:cid>')
@login_required
def imprimir_uno(cid):
    cron = Cronograma.query.get_or_404(cid)
    if cron.maestro_id != current_user.id:
        abort(403)
    return render_template('imprimir.html', cronogramas=[cron],
                            maestro=current_user)


# ─────────────────────────────────────────────
# API: validación de punteo (AJAX)
# ─────────────────────────────────────────────

@app.route('/api/validate_punteo', methods=['POST'])
@login_required
def validate_punteo():
    data  = request.get_json()
    total = int(data.get('total', 0))
    meta  = int(data.get('meta', 100))
    ok    = (total == meta)
    return jsonify({'ok': ok, 'total': total, 'meta': meta,
                    'msg': f'Total: {total} / {meta} puntos'})




# ─────────────────────────────────────────────
# CONSOLIDADO (subdirector ve todos los cronogramas)
# ─────────────────────────────────────────────

ORDEN_GRADOS = ['Primero Básico', 'Segundo Básico', 'Tercero Básico']

@app.route('/consolidado')
@login_required
def consolidado():
    """Vista consolidada: todos los cronogramas ordenados por grado y seccion."""
    todos = Cronograma.query.all()
    def sort_key(c):
        try:
            g = ORDEN_GRADOS.index(c.grado)
        except ValueError:
            g = 99
        return (g, c.seccion.upper())
    todos.sort(key=sort_key)
    return render_template('consolidado.html', cronogramas=todos,
                            maestro=current_user)


# ─────────────────────────────────────────────
# EXPORT A WORD (.docx)
# ─────────────────────────────────────────────

@app.route('/cronograma/<int:cid>/word')
@login_required
def exportar_word(cid):
    """Genera un archivo .docx con el cronograma del maestro."""
    from io import BytesIO
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        flash('Error: la librería python-docx no está instalada.', 'danger')
        return redirect(url_for('dashboard'))

    cron = Cronograma.query.get_or_404(cid)
    if cron.maestro_id != current_user.id:
        abort(403)

    doc = Document()

    # Márgenes
    for section in doc.sections:
        section.top_margin    = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin   = Cm(2)
        section.right_margin  = Cm(2)

    # Título
    titulo = doc.add_heading('', level=1)
    run = titulo.add_run(f'Cronograma de Actividades — {cron.catedra}')
    run.font.size  = Pt(14)
    run.font.bold  = True
    run.font.color.rgb = RGBColor(0x1a, 0x4b, 0x8c)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Datos generales
    doc.add_paragraph('')
    info = doc.add_paragraph()
    info.add_run('Grado: ').bold = True
    info.add_run(f'{cron.grado}   ')
    info.add_run('Sección: ').bold = True
    info.add_run(f'{cron.seccion}   ')
    info.add_run('Unidad: ').bold = True
    info.add_run(f'{cron.unidad}   ')
    info.add_run('Ciclo: ').bold = True
    info.add_run(f'{cron.ciclo}')

    info2 = doc.add_paragraph()
    info2.add_run('Período: ').bold = True
    info2.add_run(f'{cron.fecha_inicio} al {cron.fecha_fin}   ')
    info2.add_run('Punteo Meta: ').bold = True
    info2.add_run(f'{cron.punteo_meta} pts')

    info3 = doc.add_paragraph()
    info3.add_run('Maestro/a: ').bold = True
    info3.add_run(current_user.nombre)

    doc.add_paragraph('')

    # Tabla de actividades
    tabla = doc.add_table(rows=1, cols=4)
    tabla.style = 'Table Grid'
    tabla.autofit = True

    # Encabezados
    encabezados = ['#', 'Actividad', 'Aspectos a Calificar', 'Pts']
    anchos = [Cm(1), Cm(6), Cm(8), Cm(1.5)]
    hdr = tabla.rows[0].cells
    for i, (texto, ancho) in enumerate(zip(encabezados, anchos)):
        hdr[i].width = ancho
        hdr[i].text  = texto
        run = hdr[i].paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        run.font.size = Pt(10)
        # Fondo azul
        tc   = hdr[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd  = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  '1A4B8C')
        tcPr.append(shd)

    # Filas de actividades
    for act in cron.actividades:
        aspectos_txt = '\n'.join(
            f"- {a['nombre']}: {a['puntaje']} pts" for a in act.aspectos
        )
        fila = tabla.add_row().cells
        fila[0].text = str(act.numero)
        fila[1].text = act.descripcion + (f'\n{act.fecha_actividad}' if act.fecha_actividad else '')
        fila[2].text = aspectos_txt
        fila[3].text = str(act.total_actividad)
        for cell in fila:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    # Total
    doc.add_paragraph('')
    total_p = doc.add_paragraph()
    total_p.add_run('Total Unidad: ').bold = True
    total_p.add_run(f'{cron.total_unidad} / {cron.punteo_meta} pts')

    # Guardar en memoria
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    from flask import send_file
    nombre_archivo = f"cronograma_{cron.catedra}_{cron.grado}_{cron.seccion}.docx".replace(' ', '_')
    return send_file(buf, as_attachment=True,
                     download_name=nombre_archivo,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

# ─────────────────────────────────────────────
# INICIALIZACIÓN
# ─────────────────────────────────────────────

def init_db():
    with app.app_context():
        db.create_all()
        if not Maestro.query.first():
            demo = Maestro(nombre='Maestro Demo', email='demo@ineb.edu.gt')
            demo.set_password('demo1234')
            db.session.add(demo)
            db.session.commit()
            print('✓ BD creada con usuario demo: demo@ineb.edu.gt / demo1234')


init_db()

if __name__ == '__main__':
    app.run(debug=True, port=5000)
